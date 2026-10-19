"""Preserve legacy day-level evidence without fabricating article timestamps."""
from datetime import datetime, date
from pathlib import Path
import math
from openpyxl import load_workbook
from .calendar import iso
from .store import dumps


def clean(value):
    if isinstance(value, (datetime,date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def import_legacy(store, model_dir: Path | str | None = None) -> dict:
    model_dir = Path(model_dir or Path(__file__).resolve().parents[2] / "model")
    news_path = model_dir / "NaturalGasWorld_ttf_excludeWK.xlsx"
    events_path = model_dir / "news_event_impact_dataset.xlsx"
    count_news = count_events = 0
    if news_path.exists():
        wb = load_workbook(news_path,read_only=True,data_only=True)
        rows = iter(wb.active.values)
        headers = next(rows)
        with store.connection() as c:
            for values in rows:
                entry = {k: clean(v) for k,v in zip(headers,values)}
                day = str(entry["Date"])[:10]
                entry.update(timestamp_source="legacy_trading_day_aligned", timestamp_precision="date", source="Natural Gas World (legacy daily)", published_at=None, source_url=None)
                entry["Explanation_usage"] = "archival_only_not_model_input"
                c.execute("INSERT OR IGNORE INTO legacy_daily VALUES (?,?,?)", ("legacy-"+day,day,dumps(entry)))
                count_news += 1
        wb.close()
    if events_path.exists():
        wb = load_workbook(events_path,read_only=True,data_only=True)
        rows = iter(wb["event_days"].values)
        headers = next(rows)
        existing = {e["id"]:e for e in store.events(10000)}
        for values in rows:
            old = {k: clean(v) for k,v in zip(headers,values)}
            day = str(old["Date"])[:10]
            identifier = "legacy-event-"+day
            if identifier in existing:
                previous=existing[identifier]
                if previous.get('source')=='legacy_daily' and previous.get('state')=='resolved':
                    previous['state']='historical'
                    store.upsert_event(previous)
                count_events += 1
                continue
            event = {"id": identifier, "event_date": day, "title": old.get("title"), "summary": old.get("summary"),
                "category": old.get("event_category_primary"), "matched_terms": old.get("matched_event_terms"),
                "state": "historical", "known_at": iso(), "source": "legacy_daily", "source_quality": "daily_aggregate_no_original_url", "timestamp_precision": "date",
                "legacy_label": old.get("impact_persistence_label"), "legacy_outcomes": old, "outcome_available_at": None,
                "risk_flags": ["历史多条新闻日聚合；无单条来源链接与真实发布时间", "旧事后影响标签仅供档案追溯，不进入预测输入"]}
            store.upsert_event(event)
            count_events += 1
        wb.close()
    return {"legacy_daily":count_news,"events":count_events}

